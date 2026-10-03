"""
Turns an uploaded file (resume PDF/DOCX, or a JD that might be text/PDF/image)
into plain text. This is intentionally "dumb" -- it does not try to understand
the resume's structure. Structure understanding is Claude's job, in
extraction.py, because resume layouts vary too much for hand-written rules.
"""
import io
from docx import Document
import pdfplumber


class UnsupportedFileType(Exception):
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Pulls plain text out of a PDF, page by page, in reading order."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Pulls plain text out of a DOCX, including table cells (some resumes
    use tables for the header/contact-info block)."""
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts).strip()


def extract_text_from_upload(filename: str, file_bytes: bytes) -> str:
    """Dispatches by file extension. Raises UnsupportedFileType for anything
    else so the caller can return a clear error to the user instead of a
    stack trace."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
        if not text:
            raise UnsupportedFileType(
                "Could not find any selectable text in this PDF. If this is a "
                "scanned/image-based resume, please upload the original PDF/DOCX "
                "or paste the details manually."
            )
        return text
    if lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    raise UnsupportedFileType(
        f"Unsupported file type for '{filename}'. Please upload a PDF or DOCX."
    )


def is_image_file(filename: str) -> bool:
    lower = filename.lower()
    return lower.endswith((".png", ".jpg", ".jpeg", ".webp"))
